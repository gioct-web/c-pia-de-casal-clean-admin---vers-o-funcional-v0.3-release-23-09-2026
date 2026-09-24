from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/home/ubuntu/upload/EspecificaçõesdeProdutoseServiços-CasalClean.xlsx")


def parse_money(value: object) -> float:
    digits = re.sub(r"[^0-9,.-]", "", str(value)).replace(".", "").replace(",", ".")
    return float(digits or 0)


def main() -> None:
    workbook = load_workbook(SOURCE, data_only=True)
    sheet = workbook["Table 1"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    priced = []
    zero_reference = []
    for number, row in enumerate(rows, start=2):
        record = {
            "line": number,
            "product": row["Produto"],
            "type": row["Categoria/Tipo"],
            "specification": row["Especificação (Lugares/Tecido)"],
            "price": parse_money(row["Preço Sugerido (R$)"].replace("R$", "") if isinstance(row["Preço Sugerido (R$)"], str) else row["Preço Sugerido (R$)"]),
            "source": row["Fonte"],
        }
        (priced if record["price"] > 0 else zero_reference).append(record)
    print(json.dumps({"priced_rows": priced, "zero_price_reference_rows": zero_reference}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
